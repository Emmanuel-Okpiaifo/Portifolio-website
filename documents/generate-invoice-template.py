"""
Generate EDO-branded reusable invoice template (Google Docs friendly).
Run: python documents/generate-invoice-template.py
Output: documents/EDO-Invoice-Template.docx

Google Docs tips:
  - Upload to Google Drive → Open with Google Docs
  - Or File → Open in docs.google.com
  - Avoid editing complex layout in Word then re-upload; edit in Docs directly
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
LOGO_PATH = ROOT / "public" / "edo-logo.png"
OUTPUT_PATH = Path(__file__).resolve().parent / "EDO-Invoice-Template.docx"
OUTPUT_FALLBACK = Path(__file__).resolve().parent / "EDO-Invoice-Template-GDocs.docx"

CHARCOAL = RGBColor(0x1C, 0x1C, 0x1C)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLD_DARK = RGBColor(0xA8, 0x84, 0x1F)
STONE_TEXT = RGBColor(0x78, 0x71, 0x6C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PLACEHOLDER = RGBColor(0x9C, 0xA3, 0xAF)
MUTED = RGBColor(0x88, 0x88, 0x88)

# Arial + Georgia render reliably in Google Docs (no custom font install needed)
FONT_BODY = "Arial"
FONT_DISPLAY = "Georgia"


def shade_cell(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


def add_run(paragraph, text, *, bold=False, italic=False, size=10, color=CHARCOAL, font=FONT_BODY):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    return run


def gap(doc, pt=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pt)
    p.paragraph_format.space_before = Pt(0)


def section_heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    add_run(p, text, bold=True, size=9, color=GOLD)
    p.paragraph_format.space_after = Pt(4)


def simple_table(doc, rows: int, cols: int, style: str = "Table Grid"):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    return table


def build_header(doc: Document) -> None:
    """Single flat 3-column row — no nested tables."""
    table = simple_table(doc, 1, 3)
    table.columns[0].width = Cm(2.4)
    table.columns[1].width = Cm(9.6)
    table.columns[2].width = Cm(5.0)

    logo_cell, brand_cell, meta_cell = table.rows[0].cells
    for cell in (logo_cell, brand_cell, meta_cell):
        shade_cell(cell, "1C1C1C")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if LOGO_PATH.exists():
        lp = logo_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(str(LOGO_PATH), width=Cm(1.5))

    bp = brand_cell.paragraphs[0]
    add_run(bp, "EDO", bold=True, size=20, color=GOLD, font=FONT_DISPLAY)
    p2 = brand_cell.add_paragraph()
    add_run(p2, "Emmanuel (Daniel) Okpiaifo", size=9, color=RGBColor(0xCC, 0xCC, 0xCC))
    p3 = brand_cell.add_paragraph()
    add_run(p3, "Web Developer (React & WordPress) | Digital Platforms & Product Delivery", size=8, color=MUTED)

    mp = meta_cell.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(mp, "INVOICE", bold=True, size=22, color=WHITE, font=FONT_DISPLAY)

    num_p = meta_cell.add_paragraph()
    num_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(num_p, "[Invoice Number]", size=10, color=GOLD, bold=True)

    status_p = meta_cell.add_paragraph()
    status_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(status_p, "[Status: Draft / Sent / Paid]", size=8, color=RGBColor(0xD4, 0xDD, 0xD0), italic=True)

    for label, value in (
        ("Issue Date", "[DD Mon YYYY]"),
        ("Due Date", "[DD Mon YYYY]"),
        ("Currency", "NGN"),
    ):
        row_p = meta_cell.add_paragraph()
        row_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_run(row_p, f"{label}: ", size=8, color=MUTED, bold=True)
        color = PLACEHOLDER if value.startswith("[") else WHITE
        add_run(row_p, value, size=9, color=color)

    # Gold accent bar
    accent = simple_table(doc, 1, 1)
    accent.columns[0].width = Cm(17)
    shade_cell(accent.rows[0].cells[0], "C9A227")
    accent.rows[0].cells[0].text = ""
    gap(doc, 6)


def build_parties(doc: Document) -> None:
    table = simple_table(doc, 1, 2)
    table.columns[0].width = Cm(8.5)
    table.columns[1].width = Cm(8.5)
    from_cell, to_cell = table.rows[0].cells

    fp = from_cell.paragraphs[0]
    add_run(fp, "BILL FROM", bold=True, size=9, color=GOLD)
    for line, bold in (
        ("Emmanuel Daniel Okpiaifo", True),
        ("Lagos, Nigeria (on-site, hybrid or remote)", False),
        ("emmaokpiaifo@gmail.com", False),
        ("(+234) 9160852509", False),
    ):
        p = from_cell.add_paragraph()
        add_run(p, line, size=9, color=CHARCOAL if bold else STONE_TEXT, bold=bold)

    tp = to_cell.paragraphs[0]
    add_run(tp, "BILL TO", bold=True, size=9, color=GOLD)
    placeholders = [
        "[Client / Company Name]",
        "Attn: [Contact Person — optional]",
        "[Client Address — optional]",
        "[client@email.com — optional]",
        "[Client Phone — optional]",
    ]
    for ph in placeholders:
        p = to_cell.add_paragraph()
        add_run(p, ph, size=9, color=PLACEHOLDER, italic=True)

    gap(doc, 8)


def build_items(doc: Document) -> None:
    headers = ("Description", "Qty", "Rate (NGN)", "Amount (NGN)")
    rows = 1 + 5
    table = simple_table(doc, rows, 4)
    widths = (9.0, 1.8, 3.1, 3.1)
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)

    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "F5F4F0")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, title.upper(), bold=True, size=8, color=STONE_TEXT)

    line_items = [
        ("[Service / Product Name]", "[Optional description]", "1", "[0]", "[0]"),
        ("[Service / Product Name]", "[Optional description]", "1", "[0]", "[0]"),
        ("[Service / Product Name]", "[Optional description]", "1", "[0]", "[0]"),
        ("[Add row or delete unused rows]", "", "", "", ""),
        ("", "", "", "", ""),
    ]

    for r, (title, desc, qty, rate, amount) in enumerate(line_items, start=1):
        cells = table.rows[r].cells
        if r % 2 == 0:
            for c in cells:
                shade_cell(c, "FAFAF8")

        dp = cells[0].paragraphs[0]
        if title:
            add_run(dp, title, bold=True, size=9, color=PLACEHOLDER if title.startswith("[") else CHARCOAL)
        if desc:
            d2 = cells[0].add_paragraph()
            add_run(d2, desc, size=8, color=PLACEHOLDER, italic=True)

        for col, val in enumerate((qty, rate, amount), start=1):
            p = cells[col].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col > 1 else WD_ALIGN_PARAGRAPH.CENTER
            add_run(
                p,
                val,
                size=9,
                color=PLACEHOLDER if str(val).startswith("[") else CHARCOAL,
                bold=col == 3,
            )

    gap(doc, 10)


def build_payment_schedule(doc: Document) -> None:
    section_heading(doc, "PAYMENT SCHEDULE")
    table = simple_table(doc, 4, 2)
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(14.0)

    rows = [
        ("[NGN 0]", "[Milestone description — e.g. initial deposit]"),
        ("[NGN 0]", "[Milestone description]"),
        ("[NGN 0]", "[Milestone description]"),
        ("", "[Delete unused rows]"),
    ]
    for r, (amount, desc) in enumerate(rows):
        a_cell, d_cell = table.rows[r].cells
        add_run(a_cell.paragraphs[0], amount, bold=True, size=9, color=GOLD_DARK if amount else STONE_TEXT)
        add_run(
            d_cell.paragraphs[0],
            desc,
            size=9,
            color=PLACEHOLDER if desc.startswith("[") else STONE_TEXT,
            italic=desc.startswith("["),
        )

    gap(doc, 10)


def build_payment_details(doc: Document) -> None:
    section_heading(doc, "PAYMENT DETAILS")
    table = simple_table(doc, 6, 2)
    table.columns[0].width = Cm(4.5)
    table.columns[1].width = Cm(12.5)

    shade_cell(table.rows[0].cells[0], "F5F4F0")
    shade_cell(table.rows[0].cells[1], "F5F4F0")
    add_run(table.rows[0].cells[0].paragraphs[0], "Field", bold=True, size=8, color=STONE_TEXT)
    add_run(table.rows[0].cells[1].paragraphs[0], "Details", bold=True, size=8, color=STONE_TEXT)

    bank_lines = [
        ("Method", "Bank Transfer"),
        ("Bank", "First Bank Nigeria"),
        ("Account Name", "Okpiaifo Emmanuel Daniel"),
        ("Account No.", "3222647347"),
        ("Reference", "[Invoice Number]"),
    ]
    for r, (label, value) in enumerate(bank_lines, start=1):
        l_cell, v_cell = table.rows[r].cells
        add_run(l_cell.paragraphs[0], label, bold=True, size=9, color=CHARCOAL)
        add_run(
            v_cell.paragraphs[0],
            value,
            size=9,
            color=PLACEHOLDER if value.startswith("[") else STONE_TEXT,
        )

    gap(doc, 10)


def build_totals(doc: Document) -> None:
    """Spacer column + totals — flat 2-col table, not nested."""
    section_heading(doc, "TOTALS")
    table = simple_table(doc, 3, 2)
    table.columns[0].width = Cm(10.0)
    table.columns[1].width = Cm(7.0)

    total_rows = [
        ("Subtotal", "[NGN 0]"),
        ("VAT (optional — delete row if unused)", "[NGN 0]"),
        ("Total Project Cost", "[NGN 0]"),
    ]
    for r, (label, value) in enumerate(total_rows):
        l_cell, v_cell = table.rows[r].cells
        if r == 2:
            shade_cell(l_cell, "F8F3E4")
            shade_cell(v_cell, "F8F3E4")
        lp = l_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        vp = v_cell.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_run(
            lp,
            label,
            size=12 if r == 2 else 9,
            color=CHARCOAL,
            bold=r == 2,
            font=FONT_DISPLAY if r == 2 else FONT_BODY,
        )
        add_run(
            vp,
            value,
            size=13 if r == 2 else 9,
            color=GOLD_DARK if r == 2 else PLACEHOLDER,
            bold=r == 2,
            font=FONT_DISPLAY if r == 2 else FONT_BODY,
        )

    hint = doc.add_paragraph()
    hint.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(hint, "Recalculate totals after editing line items (Qty x Rate).", size=8, color=STONE_TEXT, italic=True)
    gap(doc, 10)


def build_footer(doc: Document) -> None:
    table = simple_table(doc, 1, 2)
    table.columns[0].width = Cm(8.5)
    table.columns[1].width = Cm(8.5)
    left, right = table.rows[0].cells
    shade_cell(left, "F5F4F0")
    shade_cell(right, "F5F4F0")
    add_run(left.paragraphs[0], "Thank you for your trust.", italic=True, size=11, color=CHARCOAL, font=FONT_DISPLAY)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(rp, "Built with care by Emmanuel (Daniel) Okpiaifo.", size=8, color=STONE_TEXT)


def build_instructions(doc: Document) -> None:
    gap(doc, 16)
    p = doc.add_paragraph()
    add_run(p, "— How to use (Google Docs) —", bold=True, size=12, color=CHARCOAL, font=FONT_DISPLAY)
    p.paragraph_format.space_after = Pt(8)

    tips = [
        "Upload this file to Google Drive, then right-click → Open with → Google Docs.",
        "Replace every [bracketed placeholder] with your invoice details.",
        "Gray italic text marks editable fields; bank details are pre-filled.",
        "Add or remove rows in the line-items and payment-schedule tables as needed.",
        "Update Subtotal and Total Project Cost after changing amounts.",
        "Download as PDF: File → Download → PDF Document (.pdf).",
        "Regenerate: python documents/generate-invoice-template.py",
    ]
    for tip in tips:
        bp = doc.add_paragraph(style="List Bullet")
        add_run(bp, tip, size=9, color=CHARCOAL)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(10)
    normal.font.color.rgb = CHARCOAL

    build_header(doc)
    build_parties(doc)
    build_items(doc)
    build_payment_schedule(doc)
    build_payment_details(doc)
    build_totals(doc)
    build_footer(doc)
    build_instructions(doc)

    out = OUTPUT_PATH
    try:
        doc.save(out)
    except PermissionError:
        out = OUTPUT_FALLBACK
        doc.save(out)
        print("Note: close the open template in Word/Docs, then re-run to overwrite EDO-Invoice-Template.docx")
    print(f"Created: {out}")


if __name__ == "__main__":
    main()
